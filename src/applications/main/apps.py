from decimal import Decimal, ROUND_HALF_UP
import os
from typing import List, Optional, Union
from xml.dom import minidom

from django.apps import AppConfig
from django.conf import settings
from django.db import DatabaseError

from openpyxl import load_workbook

from applications.main.d_types import DataForXML
from framework.custom_logging import logger


class MainConfig(AppConfig):
    label = "main"
    name = f"applications.{label}"


class UploadedFileObject:
    def __init__(self, file, object_id: int):
        from applications.main.models import ConstructionMaterial
        self._ConstructionMaterial = ConstructionMaterial
        self.object_id = object_id
        self.wb = load_workbook(filename=file, data_only=True)
        self.ws = self.wb.active
        self.checking_row_0 = next(self.ws.values)
        self.material_list = None

    def check_row_0(self) -> bool:
        if self.checking_row_0 != ('NAIM', 'ED_IZM', 'KOL_VO', 'PRICE', 'COST'):
            return False
        return True

    def create_obj_list(self) -> Optional[List]:
        count_obj_db = int(self._ConstructionMaterial.objects.filter(building_object__id=self.object_id).count())
        materials = []
        for n, row in enumerate(self.ws.values, start=count_obj_db):
            if row[0] is None:
                return None
            if row[2] is None:
                return None
            if row[4] is None:
                return None

            material = self._ConstructionMaterial(
                id_instance=n,
                name=row[0],
                unit=row[1] if row[1] is not None else 'ШТ',
                quantity=row[2],
                price=row[3] if row[3] is not None else Decimal('{:f}'.format(
                    (Decimal(row[4]) / Decimal(row[2])).quantize(Decimal('1.00000'), ROUND_HALF_UP).normalize())),
                total_cost=row[4],
                building_object_id=self.object_id,
            )
            materials.append(material)
        self.material_list = materials[1:]
        return self.material_list

    def save_in_db(self) -> bool:
        try:
            save = self._ConstructionMaterial.objects.bulk_create(self.material_list)
        except DatabaseError:
            logger.info("DatabaseError!")
            return False
        else:
            if save:
                return True
            return False


class FileXML:
    def __init__(self, object_id: int, object_name: str, user_id: int):
        from applications.main.models import ConstructionMaterial
        self.__ConstructionMaterial = ConstructionMaterial
        self.__object_id = object_id
        self.__object_name = object_name
        self.__user_id = user_id
        self.__xml_string = None
        self.__doc = minidom.Document()
        self.__loc_estimate = None

    def __get_data_from_base(self) -> Union[bool, List]:
        data_from_base = self.__ConstructionMaterial.objects.filter(building_object__id=self.__object_id)
        if not data_from_base:
            return False
        return data_from_base

    def __generate_data_for_xml(self) -> DataForXML:
        for value in self.__get_data_from_base():
            quantity = Decimal('{:f}'.format(Decimal(str(value.quantity)).normalize()))
            price = Decimal('{:f}'.format(Decimal(str(value.price)).normalize()))
            unit_tr = Decimal(
                '{:f}'.format((price * Decimal('0.02')).quantize(Decimal('1.00'), ROUND_HALF_UP).normalize()))
            total_tr = Decimal('{:f}'.format((unit_tr * quantity).quantize(Decimal('1.00'), ROUND_HALF_UP).normalize()))
            unit_cost = Decimal('{:f}'.format((price + unit_tr).quantize(Decimal('1.00'), ROUND_HALF_UP).normalize()))
            total_cost = Decimal('{:f}'.format(Decimal(str(value.total_cost)).normalize()))
            total_cost_total_tr = Decimal(
                '{:f}'.format((total_tr + total_cost).quantize(Decimal('1.00'), ROUND_HALF_UP).normalize())
            )
            data_modified = DataForXML(
                id_instance=str(value.id_instance),
                basis=f"СЦЕН-МТ-00{value.id_instance}" if value.id_instance <= 9 else f"СЦЕН-МТ-0{value.id_instance}"
                if value.id_instance <= 99 else f"СЦЕН-МТ-{value.id_instance}",
                name=value.name,
                unit=value.unit,
                quantity=quantity,
                price=price,
                unit_tr=unit_tr,
                total_tr=total_tr,
                unit_cost=unit_cost,
                total_cost=total_cost,
                total_cost_total_tr=total_cost_total_tr
            )
            yield data_modified

    def __generate_xml_str_header(self) -> bool:
        try:
            root = self.__doc.createElement('root')
            self.__doc.appendChild(root)

            building = self.__doc.createElement('stroyka')
            building.setAttribute('version', '1.2')
            building.setAttribute('prg', 'SXW')
            building.setAttribute('ver_prg', '')
            building.setAttribute('baza', '2017')
            building.setAttribute('idobekt', '')
            building.setAttribute('naim', '')
            building.setAttribute('nachalo', '')
            root.appendChild(building)

            obj_estimate = self.__doc.createElement('ob_smeta')
            obj_estimate.setAttribute('nomer', '1')
            obj_estimate.setAttribute('naim', self.__object_name)
            building.appendChild(obj_estimate)

            self.__loc_estimate = self.__doc.createElement('loc_smeta')
            self.__loc_estimate.setAttribute('nomer', '1')
            self.__loc_estimate.setAttribute('naim', '')
            obj_estimate.appendChild(self.__loc_estimate)

            initial_data = self.__doc.createElement('ishodnye_dannye')
            initial_data.setAttribute('zona', '3')
            initial_data.setAttribute('region', '7')
            initial_data.setAttribute('mesyac_cen', '7')
            initial_data.setAttribute('snds', '0')
            initial_data.setAttribute('vid_ohr_pp', '11')
            initial_data.setAttribute('stavka4', '6.94')
            initial_data.setAttribute('kf_zu', '0.93')
            initial_data.setAttribute('kf_vrem', '0.93')
            initial_data.setAttribute('num_post', '')
            initial_data.setAttribute('kf_p4_ohr', '1')
            initial_data.setAttribute('kf_p4_pp', '1')
            self.__loc_estimate.appendChild(initial_data)
        except Exception:
            return False
        return True

    def __generate_xml_str_body(self) -> bool:
        data_for_xml = self.__generate_data_for_xml()
        try:
            while True:
                try:
                    next_data = next(data_for_xml)
                except StopIteration:
                    break
                else:
                    valuation = self.__doc.createElement('rascenka')
                    valuation.setAttribute('npp', next_data.id_instance)
                    valuation.setAttribute('obosn', next_data.basis)
                    valuation.setAttribute('naim', next_data.name)
                    valuation.setAttribute('klv', str(next_data.quantity))
                    valuation.setAttribute('ed_izm', next_data.unit)
                    valuation.setAttribute('vid_ohr_pp', '11')  # the parameter will be fixed after import
                    valuation.setAttribute('uniq_lab_ohr_pp', '1.1')  # the parameter will be fixed after import
                    valuation.setAttribute('tip', '101')
                    valuation.setAttribute('cena_0', str(next_data.unit_cost))
                    self.__loc_estimate.appendChild(valuation)

                    price = self.__doc.createElement('cena')
                    price.setAttribute('mat', str(next_data.price))
                    price.setAttribute('tr', str(next_data.unit_tr))  # the parameter will be fixed after import
                    price.setAttribute('pryam', str(next_data.unit_cost))
                    price.setAttribute('prc_ohr', '')  # the parameter will be fixed after import
                    price.setAttribute('prc_pp', '')  # the parameter will be fixed after import
                    valuation.appendChild(price)

                    cost = self.__doc.createElement('stoimost')
                    cost.setAttribute('mat', str(next_data.total_cost))
                    cost.setAttribute('tr', str(next_data.total_tr))  # the parameter will be fixed after import
                    cost.setAttribute('pryam', str(next_data.total_cost_total_tr))
                    cost.setAttribute('ohr', '0')
                    cost.setAttribute('pp', '0')
                    valuation.appendChild(cost)

                    materials = self.__doc.createElement('materialy')
                    valuation.appendChild(materials)

                    resources = self.__doc.createElement('resurs')
                    resources.setAttribute('obosn', next_data.basis)
                    resources.setAttribute('kodcic', '')
                    resources.setAttribute('naim', next_data.name)
                    resources.setAttribute('ed_izm', next_data.unit)
                    resources.setAttribute('norma', '1')
                    resources.setAttribute('klv', str(next_data.quantity))
                    resources.setAttribute('cena_0', str(next_data.price))
                    resources.setAttribute('stm_0', str(next_data.total_cost))
                    resources.setAttribute('tr', str(next_data.unit_tr))
                    resources.setAttribute('tr_rub', '0')
                    resources.setAttribute('cena_tr', '2')  # the parameter will be fixed after import
                    materials.appendChild(resources)
        except Exception:
            return False
        return True

    def __collect_xml_string(self) -> bool:
        try:
            xml_string = self.__doc.toprettyxml(indent=" ")
            self.__xml_string = xml_string[:-8]
        except Exception:
            return False
        return True

    def __add_xml_str_footer(self) -> bool:
        file_path_footer_part = f"{settings.MEDIA_ROOT}/_footer_part.txt"
        try:
            with open(f"{file_path_footer_part}", "r") as f_footer_part:
                self.__xml_string += f_footer_part.read()
        except Exception:
            return False
        return True

    def save_xml_file(self) -> bool:
        if not self.__get_data_from_base():
            logger.debug(f"No data on the object!")
            return False
        if not self.__generate_xml_str_header():
            logger.debug(f"Error: header not generated!")
            return False
        if not self.__generate_xml_str_body():
            logger.debug(f"Error: body not generated!")
            return False
        if not self.__collect_xml_string():
            logger.debug(f"Error: xml-string not collected!")
            return False
        if not self.__add_xml_str_footer():
            logger.debug(f"Error: no footer added!")
            return False

        path = f"{settings.MEDIA_ROOT}/{self.__user_id}/xml/"
        file_path_xml = f"{settings.MEDIA_ROOT}/{self.__user_id}/xml/Materials.xml"

        try:
            if not os.path.exists(path):
                os.makedirs(path)
        except Exception:
            logger.debug(f"Directory creation error!")
            return False

        try:
            with open(f"{file_path_xml}", "w") as f_basic:
                f_basic.write(self.__xml_string)
        except Exception:
            logger.debug(f"File save error!")
            return False

        logger.debug(f"File saved successfully!")
        return True
